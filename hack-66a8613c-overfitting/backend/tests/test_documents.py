from datetime import datetime, timezone
import os
from pathlib import Path
import subprocess
import unittest
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZipFile
from xml.etree import ElementTree

from docx import Document

from backend.app.errors import LocalError
from backend.app.exporting.documents import ProtocolExporter
from backend.app.schemas import PreparationResult
from backend.app.sources.storage import MeetingStorage
from backend.tests.helpers import temporary_directory


def approved_result(meeting_id):
    return PreparationResult.model_validate({
        'meeting': {'id': meeting_id, 'title': '../../Жоба <&> \x00',
                    'meeting_date': '2026-09-23', 'timezone': 'Asia/Qyzylorda',
                    'approved_at': datetime.now(timezone.utc)},
        'job': {'id': 'job', 'meeting_id': meeting_id, 'status': 'ready'},
        'analysis_completed': True, 'summary': 'Қазақша: ә ғ қ ң ө ұ ү һ і. Итог.',
        'speakers': [{'id': 's', 'meeting_id': meeting_id, 'label': 'speaker_1', 'name': 'Имя'}],
        'utterances': [{'id': 'u', 'meeting_id': meeting_id, 'speaker_id': 's',
                        'start_seconds': 61, 'end_seconds': 65, 'text': '<script>&Әңгіме\x01'}],
        'key_points': [{'id': 'k', 'meeting_id': meeting_id, 'position': 1, 'direction': 'Доклад',
                        'metric': 'не указан', 'problem': 'Мәселе', 'source_utterance_ids': ['u']}],
        'topics': [{'id': 't', 'meeting_id': meeting_id, 'position': 1, 'title': 'Жоба',
                    'summary': 'Описание темы', 'source_utterance_ids': ['u']}],
        'action_items': [{'id': 'a', 'meeting_id': meeting_id, 'topic_id': 't', 'text': 'Подготовить ' + 'Ә' * 1000,
                          'responsible': 'не указан', 'deadline_original': 'в пятницу', 'utterance_id': 'u',
                          'timestamp_seconds': 61, 'requires_review': True, 'source_utterance_ids': ['u']}],
    })


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = temporary_directory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.storage = MeetingStorage(self.root / 'data')
        self.id = str(uuid4())
        self.directory = self.storage.directory(self.id)
        self.directory.mkdir()
        self.result = approved_result(self.id)
        self.exporter = ProtocolExporter(self.storage, 'local-soffice')

    def test_docx_structure_unicode_plain_text_and_sources(self):
        path = self.exporter.generate(self.result, 'docx')
        self.assertEqual(path.parent, self.directory)
        self.assertRegex(path.name, r'^export-[0-9a-f-]+\.docx$')
        document = Document(path)
        text = '\n'.join(p.text for p in document.paragraphs)
        for expected in ['Саммари по ключевым пунктам', 'Поручения', 'Тема 1 -- Жоба',
                         'Приложение — транскрипт', 'ә ғ қ ң ө ұ ү һ і', '01:01.0',
                         'Имя', '<script>&Әңгіме', 'Проверено и утверждено человеком']:
            self.assertIn(expected, text)
        self.assertNotIn('\x00', text)
        self.assertNotIn('\x01', text)
        self.assertEqual(len(document.tables), 3)
        self.assertEqual([c.text for c in document.tables[0].rows[0].cells],
                         ['Направление / доклад', 'Показатель', 'Проблема'])
        self.assertEqual([c.text for c in document.tables[1].rows[0].cells],
                         ['Поручение', 'Ответственный', 'Срок'])
        action_text = document.tables[1].cell(1, 0).text
        self.assertIn('Ә' * 1000, action_text)
        self.assertIn('Требует проверки', action_text)
        self.assertIn('01:01.0 [u]', action_text)
        self.assertEqual(str(document.styles['Heading 1'].font.color.rgb), '245A81')
        self.assertEqual(document.styles['Normal'].font.name, 'DejaVu Sans')
        with ZipFile(path) as archive:
            xml = archive.read('word/document.xml').decode()
            self.assertIn('E8EFF7', xml)
            self.assertIn('&lt;script&gt;&amp;', xml)
            for name in archive.namelist():
                if name.endswith('.rels'):
                    root = ElementTree.fromstring(archive.read(name))
                    self.assertFalse(any(e.get('TargetMode') == 'External' for e in root))
        self.assertEqual(list(self.directory.iterdir()), [path])

    def test_requires_approval_ready_and_analysis(self):
        for field in ('approval', 'status', 'analysis'):
            result = self.result.model_copy(deep=True)
            if field == 'approval':
                result.meeting.approved_at = None
            elif field == 'status':
                result.job.status = 'analyzing'
            else:
                result.analysis_completed = False
            with self.subTest(field=field), self.assertRaises(LocalError) as caught:
                self.exporter.generate(result, 'docx')
            self.assertEqual(caught.exception.code, 'not_approved')
            self.assertEqual(caught.exception.status, 409)
        self.assertEqual(list(self.directory.iterdir()), [])

    def test_empty_protocol_does_not_invent_action_rows(self):
        self.result.action_items = []
        self.result.topics = []
        document = Document(self.exporter.generate(self.result, 'docx'))
        self.assertEqual(len(document.tables[1].rows), 1)

    def pdf_run(self, command, **options):
        directory = Path(options['cwd'])
        self.assertTrue(directory.is_relative_to(self.directory))
        self.assertFalse(options['shell'])
        self.assertEqual(options['timeout'], 60)
        self.assertEqual(options['stdout'], subprocess.DEVNULL)
        self.assertEqual(options['stderr'], subprocess.DEVNULL)
        for key in ('HOME', 'TMPDIR', 'TMP', 'TEMP', 'XDG_CACHE_HOME', 'XDG_CONFIG_HOME', 'XDG_DATA_HOME'):
            self.assertEqual(options['env'][key], str(directory))
        self.assertIn('-env:UserInstallation=' + (directory / 'profile').as_uri(), command)
        self.assertIn('pdf:writer_pdf_Export', command)
        (directory / 'protocol.pdf').write_bytes(b'%PDF-1.7\nunit-test-placeholder')
        return subprocess.CompletedProcess(command, 0)

    def test_pdf_mock_local_conversion_and_cleanup(self):
        with patch('backend.app.exporting.documents.shutil.which', return_value='local-soffice'), \
                patch('backend.app.exporting.documents.subprocess.run', side_effect=self.pdf_run) as run:
            path = self.exporter.generate(self.result, 'pdf')
        self.assertTrue(path.read_bytes().startswith(b'%PDF-'))
        self.assertEqual(path.parent, self.directory)
        self.assertEqual(list(self.directory.iterdir()), [path])
        run.assert_called_once()

    def test_missing_libreoffice(self):
        with patch('backend.app.exporting.documents.shutil.which', return_value=None), \
                patch('backend.app.exporting.documents.subprocess.run') as run, self.assertRaises(LocalError) as caught:
            self.exporter.generate(self.result, 'pdf')
        self.assertEqual(caught.exception.code, 'libreoffice_missing')
        run.assert_not_called()
        self.assertEqual(list(self.directory.iterdir()), [])

    def test_pdf_failures_are_safe_and_leave_no_partial_file(self):
        def incomplete(command, **kwargs):
            (Path(kwargs['cwd']) / 'protocol.pdf').write_bytes(b'partial private data')
            return subprocess.CompletedProcess(command, 0)
        failures = [
            (subprocess.TimeoutExpired('PRIVATE PATH', 60), 'export_timeout'),
            (OSError('PRIVATE PATH'), 'pdf_conversion_failed'),
            (lambda *a, **k: subprocess.CompletedProcess(a[0], 1), 'pdf_conversion_failed'),
            (lambda *a, **k: subprocess.CompletedProcess(a[0], 0), 'pdf_conversion_failed'),
            (incomplete, 'pdf_conversion_failed'),
        ]
        for effect, expected in failures:
            with self.subTest(expected=expected), patch('backend.app.exporting.documents.shutil.which', return_value='local'), \
                    patch('backend.app.exporting.documents.subprocess.run', side_effect=effect), self.assertRaises(LocalError) as caught:
                self.exporter.generate(self.result, 'pdf')
            self.assertEqual(caught.exception.code, expected)
            self.assertNotIn('PRIVATE', str(caught.exception))
            self.assertEqual(list(self.directory.iterdir()), [])

    def test_size_write_and_atomic_publish_errors(self):
        for target, options in [
            ('backend.app.exporting.documents.MAX_EXPORT_BYTES', {'new': 1}),
            ('backend.app.exporting.documents.write_docx', {'side_effect': OSError('PRIVATE')}),
            ('backend.app.exporting.documents.os.replace', {'side_effect': OSError('PRIVATE')}),
        ]:
            with self.subTest(target=target), patch(target, **options), self.assertRaises(LocalError) as caught:
                self.exporter.generate(self.result, 'docx')
            self.assertEqual(caught.exception.code, 'export_failed')
            self.assertNotIn('PRIVATE', str(caught.exception))
            self.assertEqual(list(self.directory.iterdir()), [])

    def test_unsafe_id_and_reparse_point_are_rejected(self):
        self.result.meeting.id = '../outside'
        with self.assertRaises(LocalError):
            self.exporter.generate(self.result, 'docx')
        self.result.meeting.id = self.id
        original = Path.is_junction
        with patch.object(Path, 'is_junction', lambda p: p == self.directory or original(p)), self.assertRaises(LocalError) as caught:
            self.exporter.generate(self.result, 'docx')
        self.assertEqual(caught.exception.code, 'unsafe_path')
        self.assertEqual(list(self.directory.iterdir()), [])


if __name__ == '__main__':
    unittest.main()
