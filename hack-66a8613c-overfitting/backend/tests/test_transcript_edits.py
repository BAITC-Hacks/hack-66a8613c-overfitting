from io import BytesIO
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

from backend.app.database import connect
from backend.tests import test_review as fixtures


class TranscriptEditsTests(unittest.TestCase):
    setUp = fixtures.ReviewTests.setUp
    approve = fixtures.ReviewTests.approve

    def test_reassign_to_another_existing_speaker_in_same_meeting(self):
        with connect(self.service.repository.database) as db:
            db.execute('INSERT INTO speakers VALUES (?,?,?,?)', ('second-speaker', self.mid, 'speaker_2', 'Спикер 2'))
        self.approve()
        body = self.client.patch(self.url, json={'utterances': [{'id': self.result.utterances[0].id, 'speaker_id': 'second-speaker'}]}).json()
        self.assertEqual(body['utterances'][0]['speaker_id'], 'second-speaker')
        self.assertIsNone(body['meeting']['approved_at'])
        self.assertEqual(body['utterances'][0]['text'], self.result.utterances[0].text)
    def test_edit_unassign_reassign_and_export(self):
        self.approve()
        uid = self.result.utterances[0].id
        response = self.client.patch(self.url, json={'utterances': [{'id': uid, 'text': 'Исправленный мәтін', 'speaker_id': None}]})
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertIsNone(body['meeting']['approved_at'])
        self.assertIsNone(body['utterances'][0]['speaker_id'])
        self.assertEqual(body['summary'], self.result.summary)
        self.assertTrue(body['action_items'][0]['requires_review'])
        self.client.patch(self.url, json={'utterances': [{'id': uid, 'speaker_id': self.speaker.id}],
                                        'speakers': [{'id': self.speaker.id, 'name': 'Әлия'}]})
        self.approve()
        response = self.client.get(self.url + '/exports/docx')
        text = '\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs)
        self.assertIn('Исправленный мәтін', text)
        self.assertIn('Әлия', text)
        def convert(command, **kwargs):
            paragraphs = '\n'.join(p.text for p in Document(command[-1]).paragraphs)
            self.assertIn('Исправленный мәтін', paragraphs)
            self.assertIn('Әлия', paragraphs)
            (Path(kwargs['cwd']) / 'protocol.pdf').write_bytes(b'%PDF-test')
            return SimpleNamespace(returncode=0)
        with patch('backend.app.exporting.documents.shutil.which', return_value='mock'), patch('backend.app.exporting.documents.subprocess.run', side_effect=convert):
            self.assertEqual(self.client.get(self.url + '/exports/pdf').status_code, 200)

    def test_invalid_and_foreign_entities_rollback(self):
        uid = self.result.utterances[0].id
        other = fixtures.ready_meeting(self.service)
        for edit in ({'id': 'missing', 'text': 'Changed'}, {'id': other.utterances[0].id, 'text': 'Changed'},
                     {'id': uid, 'speaker_id': other.speakers[0].id}, {'id': uid, 'speaker_id': 'missing'}):
            response = self.client.patch(self.url, json={'utterances': [{'id': uid, 'text': 'Rollback'}, edit]})
            # Duplicate IDs are validation errors; test foreign speaker separately below.
            self.assertIn(response.status_code, (404, 422))
            self.assertEqual(self.service.repository.result(self.mid).utterances[0].text, self.result.utterances[0].text)
            self.assertEqual(self.client.patch(self.url, json={'utterances': [edit]}).status_code, 404)
        for edit in ({'id': uid}, {'id': uid, 'text': None}, {'id': uid, 'speaker_id': 42},
                     {'id': uid, 'text': ''}, {'id': uid, 'text': 'x' * 8001}, {'id': uid, 'start_seconds': 2}):
            self.assertEqual(self.client.patch(self.url, json={'utterances': [edit]}).status_code, 422)

    def test_deadline_saved_and_recalculated_without_inventing(self):
        self.assertEqual(self.result.action_items[0].deadline_date.isoformat(), '2026-09-24')
        for phrase, expected, review in [('  через 2 недели  ', '2026-10-07', False), ('непонятно когда', None, True)]:
            body = self.client.patch(self.url, json={'action_items': [{'id': self.action.id, 'deadline_original': phrase, 'requires_review': False}]}).json()
            action = body['action_items'][0]
            self.assertEqual(action['deadline_original'], phrase)
            self.assertEqual(action['deadline_date'], expected)
            self.assertEqual(action['requires_review'], review)
        with connect(self.service.repository.database) as db:
            db.execute('UPDATE meetings SET timezone=NULL WHERE id=?', (self.mid,))
        body = self.client.patch(self.url, json={'action_items': [{'id': self.action.id, 'deadline_original': 'завтра', 'requires_review': False}]}).json()
        self.assertIsNone(body['action_items'][0]['deadline_date'])
        self.assertTrue(body['action_items'][0]['requires_review'])
