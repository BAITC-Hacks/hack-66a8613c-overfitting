"""Build approved protocols locally; user text is always plain document text."""
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from uuid import uuid4

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from ..errors import LocalError
from ..schemas import PreparationResult
from ..sources.storage import MeetingStorage, confined

MAX_EXPORT_BYTES = 32 * 1024 * 1024


def clean(value) -> str:
    """Retain XML 1.0 Unicode characters, never parse supplied markup."""
    return ''.join(char for char in str(value) if char in '\t\n\r' or
                   0x20 <= ord(char) <= 0xD7FF or 0xE000 <= ord(char) <= 0xFFFD or
                   0x10000 <= ord(char) <= 0x10FFFF)


def timestamp(seconds: float) -> str:
    minutes, seconds = divmod(seconds, 60)
    return f'{int(minutes):02d}:{seconds:04.1f}'


def table(document, headings, rows):
    result = document.add_table(rows=1, cols=len(headings))
    result.style = 'Table Grid'
    result.autofit = False
    for cell, heading in zip(result.rows[0].cells, headings):
        cell.text = clean(heading)
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'E8EFF7')
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement('w:tblHeader')
    result.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        for cell, value in zip(result.add_row().cells, values):
            cell.text = clean(value)
            # Permit word wrapping, including long unbroken user text; no fixed row heights.
            for paragraph in cell.paragraphs:
                wrap = OxmlElement('w:wordWrap')
                wrap.set(qn('w:val'), '0')
                paragraph._p.get_or_add_pPr().append(wrap)
    return result


def write_docx(result: PreparationResult, path: Path):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(20)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = document.styles[name]
        style.font.name = 'DejaVu Sans'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'DejaVu Sans')
        if name != 'Normal':
            style.font.color.rgb = RGBColor.from_string('245A81')
    document.styles['Normal'].font.size = Pt(10)
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    document.core_properties.title = ''
    document.add_heading(clean(result.meeting.title), 0)
    document.add_paragraph(f'Дата: {result.meeting.meeting_date or "не указана"}; '
                           f'часовой пояс: {clean(result.meeting.timezone or "не указан")}')
    names = list(dict.fromkeys(result.meeting.participants + [s.name for s in result.speakers
                  if s.name and not re.fullmatch(r'Спикер\s+\d+', s.name)]))
    document.add_paragraph('Известные участники: ' + clean(', '.join(names) or 'не указаны'))
    document.add_paragraph('Машинная подготовка: ' + ('да' if result.meeting.machine_prepared else 'нет'))
    document.add_paragraph('Проверено и утверждено человеком: ' + result.meeting.approved_at.isoformat())
    document.add_heading('Саммари по ключевым пунктам', 1)
    document.add_paragraph(clean(result.summary))
    table(document, ['Направление / доклад', 'Показатель', 'Проблема'],
          [(k.direction, k.metric, k.problem) for k in result.key_points])
    sources = {u.id: u for u in result.utterances}

    def actions(items):
        rows = []
        for item in items:
            ids = item.source_utterance_ids or [item.utterance_id]
            links = [f'{timestamp(sources[i].start_seconds)} [{i}]' for i in ids if i in sources]
            text = item.text
            if item.requires_review:
                text += '\nТребует проверки'
            if links:
                text += '\nИсточники: ' + '; '.join(links)
            rows.append((text, item.responsible, item.deadline_original))
        table(document, ['Поручение', 'Ответственный', 'Срок'], rows)

    document.add_heading('Поручения', 1)
    actions(result.action_items)
    for topic in sorted(result.topics, key=lambda t: t.position):
        document.add_heading(clean(f'Тема {topic.position} -- {topic.title}'), 1)
        document.add_paragraph(clean(topic.summary))
        actions([a for a in result.action_items if a.topic_id == topic.id])
    document.add_heading('Приложение — транскрипт', 1)
    speakers = {s.id: s.name or s.label for s in result.speakers}
    for utterance in result.utterances:
        paragraph = document.add_paragraph()
        paragraph.add_run(clean(f'{timestamp(utterance.start_seconds)}–{timestamp(utterance.end_seconds)} '
                         f'{speakers.get(utterance.speaker_id, "Говорящий не определён")} '
                         f'[{utterance.id}]\n')).bold = True
        paragraph.add_run(clean(utterance.text))
        if utterance.requires_review:
            paragraph.add_run('\nТребует проверки').italic = True
    document.save(path)


class ProtocolExporter:
    def __init__(self, storage: MeetingStorage, libreoffice: str):
        self.storage = storage
        self.libreoffice = libreoffice

    def _check_file(self, path: Path, code: str):
        path = confined(self.storage.root, path)
        if not path.is_file() or not 0 < path.stat().st_size <= MAX_EXPORT_BYTES:
            raise LocalError(code, 500)

    def _pdf(self, directory: Path, docx: Path) -> Path:
        executable = shutil.which(self.libreoffice)
        if not executable:
            raise LocalError('libreoffice_missing', 503)
        profile = confined(self.storage.root, directory / 'profile')
        profile.mkdir()
        environment = os.environ.copy()
        for key in ('HOME', 'TMPDIR', 'TMP', 'TEMP', 'XDG_CACHE_HOME', 'XDG_CONFIG_HOME', 'XDG_DATA_HOME'):
            environment[key] = str(directory)
        command = [executable, '-env:UserInstallation=' + profile.as_uri(), '--headless',
                   '--nologo', '--nodefault', '--nofirststartwizard', '--norestore',
                   '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(directory), str(docx)]
        try:
            completed = subprocess.run(command, cwd=directory, env=environment, stdin=subprocess.DEVNULL,
                                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60,
                                       creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0,
                                       check=False, shell=False)
        except subprocess.TimeoutExpired:
            raise LocalError('export_timeout', 504) from None
        except FileNotFoundError:
            raise LocalError('libreoffice_missing', 503) from None
        except OSError:
            raise LocalError('pdf_conversion_failed', 500) from None
        if completed.returncode:
            raise LocalError('pdf_conversion_failed', 500)
        pdf = confined(self.storage.root, directory / 'protocol.pdf')
        self._check_file(pdf, 'pdf_conversion_failed')
        with pdf.open('rb') as stream:
            if stream.read(5) != b'%PDF-':
                raise LocalError('pdf_conversion_failed', 500)
        return pdf

    def generate(self, result: PreparationResult, format: str) -> Path:
        if not result.meeting.approved_at or not result.analysis_completed or result.job.status != 'ready':
            raise LocalError('not_approved', 409)
        if format not in ('docx', 'pdf'):
            raise LocalError('export_failed', 422)
        destination = None
        try:
            directory = self.storage.directory(result.meeting.id)
            with tempfile.TemporaryDirectory(prefix='.export-', dir=directory) as temporary:
                work = confined(self.storage.root, Path(temporary))
                docx = confined(self.storage.root, work / 'protocol.docx')
                write_docx(result, docx)
                self._check_file(docx, 'export_failed')
                output = self._pdf(work, docx) if format == 'pdf' else docx
                self._check_file(output, 'export_failed')
                destination = self.storage.file(result.meeting.id, f'export-{uuid4()}.{format}')
                os.replace(output, destination)
            return destination
        except Exception as error:
            if destination is not None:
                try:
                    confined(self.storage.root, destination).unlink(missing_ok=True)
                except (OSError, LocalError):
                    pass
            if isinstance(error, LocalError):
                raise
            raise LocalError('export_failed', 500) from None
