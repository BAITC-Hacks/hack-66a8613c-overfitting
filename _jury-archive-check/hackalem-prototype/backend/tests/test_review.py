from io import BytesIO
from pathlib import Path
import sqlite3
import subprocess
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient

from backend.app import config
from backend.app.database import connect
from backend.app.main import app
from backend.app.processing.analysis import AnalysisDraft
from backend.app.schemas import Meeting, Speaker, Utterance
from backend.tests.helpers import temporary_directory


def ready_meeting(service):
    mid = str(uuid4())
    meeting = Meeting(id=mid, title='../../Жоба <&>\r\nPRIVATE', meeting_date='2026-09-23', timezone='Asia/Qyzylorda')
    directory = service.storage.directory(mid)
    directory.mkdir()
    source = service.storage.file(mid, 'source.wav')
    source.write_bytes(b'unit fixture; not audio')
    service.repository.create(meeting, source, source.stat().st_size, 'fixture.wav')
    speaker = Speaker(id=str(uuid4()), meeting_id=mid, label='speaker_1', name='Спикер 1')
    utterance = Utterance(id=str(uuid4()), meeting_id=mid, speaker_id=speaker.id, start_seconds=1,
                          end_seconds=3, text='Подготовить отчёт. Срок завтра.')
    service.repository.save_transcript(mid, [speaker], [utterance], 'ru')
    service.repository.save_analysis(mid, AnalysisDraft.model_validate({
        'summary': 'Обсудили отчёт.', 'key_points': [], 'topics': [
            {'title': 'Отчёт', 'summary': 'Подготовка отчёта', 'source_utterance_ids': [utterance.id],
             'requires_review': True, 'action_items': [
                 {'text': 'Подготовить отчёт', 'responsible': 'не указан', 'deadline_original': 'завтра',
                  'requires_review': True, 'source_utterance_ids': [utterance.id]}]}]}))
    return service.repository.result(mid)


class ReviewTests(unittest.TestCase):
    def setUp(self):
        temp = temporary_directory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.enterContext(patch.multiple(config, DATA_DIR=self.root, DATABASE_PATH=self.root / 'test.sqlite3'))
        self.client = self.enterContext(TestClient(app))
        self.service = app.state.service
        self.result = ready_meeting(self.service)
        self.mid = self.result.meeting.id
        self.url = f'/api/meetings/{self.mid}'
        self.action = self.result.action_items[0]
        self.speaker = self.result.speakers[0]

    def approve(self):
        response = self.client.patch(self.url, json={'action_items': [{'id': self.action.id, 'requires_review': False}], 'approve': True})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_patch_fields_preserves_sources_and_returns_fresh_result(self):
        response = self.client.patch(self.url, json={
            'speakers': [{'id': self.speaker.id, 'name': '  Әлия <&>  '}],
            'action_items': [{'id': self.action.id, 'text': 'Проверить отчёт', 'responsible': 'Әлия',
                              'deadline_original': 'не указан', 'requires_review': False}]})
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body['speakers'][0]['name'], 'Әлия <&>')
        self.assertEqual(body['meeting']['participants'], ['Әлия <&>'])
        action = body['action_items'][0]
        self.assertEqual(action['text'], 'Проверить отчёт')
        self.assertEqual(action['source_utterance_ids'], self.action.source_utterance_ids)
        self.assertEqual(action['utterance_id'], self.action.utterance_id)
        self.assertEqual(action['timestamp_seconds'], 1)
        self.assertIsNone(action['deadline_date'])
        self.assertTrue(action['requires_review'])
        self.assertEqual(self.client.get(self.url + '/result').json(), body)

    def test_validation_rejects_unknown_immutable_null_empty_and_wrong_types(self):
        invalid = [
            {'speakers': [{'id': self.speaker.id, 'name': '  '}]},
            {'speakers': [{'id': self.speaker.id, 'name': 'x' * 201}]},
            {'speakers': [{'id': self.speaker.id, 'name': 'PRIVATE\u0000'}]},
            {'speakers': [{'id': self.speaker.id, 'name': None}]},
            {'speakers': [{'id': self.speaker.id, 'name': 'A', 'meeting_id': 'other'}]},
            {'action_items': [{'id': self.action.id}]},
            {'action_items': [{'id': self.action.id, 'requires_review': 'false'}]},
            {'action_items': [{'id': self.action.id, 'responsible': None}]},
            {'action_items': [{'id': self.action.id, 'text': ''}]},
            {'action_items': [{'id': self.action.id, 'text': 'x' * 8001}]},
            {'action_items': [{'id': self.action.id, 'deadline_original': 'x' * 501}]},
            {'action_items': [{'id': self.action.id, 'source_utterance_ids': ['other']}]},
            {'speakers': [{'id': self.speaker.id, 'name': 'A'}] * 2},
            {'action_items': [{'id': self.action.id, 'text': 'PRIVATE'}] * 2},
            {'approve': 'true'}, {'approved_at': '2030-01-01'}, {'utterances': [{'id': 'u', 'text': None}]},
        ]
        for body in invalid:
            with self.subTest(body=body):
                response = self.client.patch(self.url, json=body)
                self.assertEqual(response.status_code, 422)
                self.assertEqual(response.json()['detail']['code'], 'invalid_edits')
                self.assertNotIn('PRIVATE', response.text)
        self.assertEqual(self.client.get(self.url + '/result').json()['action_items'][0]['text'], self.action.text)

    def test_missing_and_foreign_entities_are_404_and_rollback(self):
        other = ready_meeting(self.service)
        for edits in [
            {'speakers': [{'id': self.speaker.id, 'name': 'Changed'}, {'id': other.speakers[0].id, 'name': 'Other'}]},
            {'action_items': [{'id': self.action.id, 'text': 'Changed'}, {'id': 'missing', 'text': 'Other'}]},
        ]:
            response = self.client.patch(self.url, json=edits)
            self.assertEqual(response.status_code, 404)
        saved = self.service.repository.result(self.mid)
        self.assertEqual(saved.speakers[0].name, self.speaker.name)
        self.assertEqual(saved.action_items[0].text, self.action.text)
        self.assertEqual(self.client.patch('/api/meetings/missing', json={}).status_code, 404)
        for format in ('docx', 'pdf'):
            self.assertEqual(self.client.get(f'/api/meetings/missing/exports/{format}').status_code, 404)

    def test_approval_requires_ready_analysis_and_resolved_actions(self):
        response = self.client.patch(self.url, json={'approve': True})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()['detail']['code'], 'review_required')
        self.service.repository.update(self.mid, 'analyzing')
        self.assertEqual(self.client.patch(self.url, json={}).status_code, 409)
        self.service.repository.update(self.mid, 'ready')
        with connect(self.service.repository.database) as db:
            db.execute('DELETE FROM analyses WHERE meeting_id=?', (self.mid,))
        self.assertEqual(self.client.patch(self.url, json={'approve': True}).status_code, 409)

    def test_approval_is_saved_and_edits_revoke_it(self):
        approved = self.approve()
        self.assertIsNotNone(approved['meeting']['approved_at'])
        self.assertFalse(approved['requires_review'])
        self.assertFalse(approved['topics'][0]['requires_review'])
        again = self.client.patch(self.url, json={'approve': True}).json()
        self.assertEqual(approved['meeting']['approved_at'], again['meeting']['approved_at'])
        for edits in [{'speakers': [{'id': self.speaker.id, 'name': 'Имя'}]},
                      {'action_items': [{'id': self.action.id, 'deadline_original': 'в пятницу'}]}]:
            response = self.client.patch(self.url, json=edits)
            self.assertIsNone(response.json()['meeting']['approved_at'])
            self.assertTrue(response.json()['requires_review'])
            self.assertEqual(self.client.get(self.url + '/exports/docx').status_code, 409)
            self.approve()

    def test_exports_require_approval_before_any_file_or_process(self):
        with patch('backend.app.exporting.documents.ProtocolExporter.generate') as generate:
            for format in ('docx', 'pdf'):
                response = self.client.get(self.url + '/exports/' + format)
                self.assertEqual(response.status_code, 409)
                self.assertEqual(response.json()['detail']['code'], 'not_approved')
        generate.assert_not_called()
        self.assertEqual(list((self.root / self.mid).iterdir()), [self.root / self.mid / 'source.wav'])

    def test_docx_download_headers_saved_file_and_updated_contents(self):
        self.client.patch(self.url, json={'speakers': [{'id': self.speaker.id, 'name': 'Әлия'}]})
        self.approve()
        response = self.client.get(self.url + '/exports/docx')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['content-type'], 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        self.assertEqual(response.headers['content-disposition'], f'attachment; filename="meeting-{self.mid}.docx"')
        self.assertEqual(response.headers['cache-control'], 'no-store')
        self.assertEqual(response.headers['x-content-type-options'], 'nosniff')
        self.assertNotIn('PRIVATE', str(response.headers))
        document = Document(BytesIO(response.content))
        self.assertIn('Әлия', '\n'.join(p.text for p in document.paragraphs))
        self.assertGreaterEqual(len(document.tables), 3)
        with connect(self.service.repository.database) as db:
            record = db.execute('SELECT * FROM exports WHERE meeting_id=?', (self.mid,)).fetchone()
            path = Path(record['storage_path'])
            self.assertEqual(path.parent, self.root / self.mid)
            self.assertEqual(path.read_bytes(), response.content)
        self.client.patch(self.url, json={'action_items': [{'id': self.action.id, 'text': 'Правка'}]})
        self.assertEqual(self.client.get(self.url + '/exports/docx').status_code, 409)
        with connect(self.service.repository.database) as db:
            self.assertEqual(db.execute('SELECT count(*) FROM exports').fetchone()[0], 0)

    def test_pdf_http_headers_with_mocked_local_converter(self):
        self.approve()
        def convert(command, **kwargs):
            (Path(kwargs['cwd']) / 'protocol.pdf').write_bytes(b'%PDF-1.7\nunit fixture')
            return SimpleNamespace(returncode=0)
        with patch('backend.app.exporting.documents.shutil.which', return_value='local-soffice'), \
             patch('backend.app.exporting.documents.subprocess.run', side_effect=convert):
            response = self.client.get(self.url + '/exports/pdf')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['content-type'], 'application/pdf')
        self.assertEqual(response.headers['content-disposition'], f'attachment; filename="meeting-{self.mid}.pdf"')
        self.assertTrue(response.content.startswith(b'%PDF-'))
        self.assertFalse(list((self.root / self.mid).glob('.export-*')))

    def test_pdf_safe_errors_without_partial_exports(self):
        self.approve()
        with patch('backend.app.exporting.documents.shutil.which', return_value=None):
            response = self.client.get(self.url + '/exports/pdf')
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.json()['detail']['code'], 'libreoffice_missing')
        for failure, code, status in [(OSError('PRIVATE'), 'pdf_conversion_failed', 500),
                                      (subprocess.TimeoutExpired('PRIVATE', 60), 'export_timeout', 504)]:
            with patch('backend.app.exporting.documents.shutil.which', return_value='local-soffice'), \
                 patch('backend.app.exporting.documents.subprocess.run', side_effect=failure):
                response = self.client.get(self.url + '/exports/pdf')
            self.assertEqual(response.status_code, status)
            self.assertEqual(response.json()['detail']['code'], code)
            self.assertNotIn('PRIVATE', response.text)
        self.assertFalse(list((self.root / self.mid).glob('*export*')))

    def test_export_db_failure_removes_artifact_and_is_safe(self):
        self.approve()
        with patch.object(self.service.repository, 'record_export', side_effect=sqlite3.OperationalError('PRIVATE')):
            response = self.client.get(self.url + '/exports/docx')
        self.assertEqual(response.status_code, 500)
        self.assertEqual(response.json()['detail']['code'], 'export_failed')
        self.assertFalse(list((self.root / self.mid).glob('*export*')))

    def test_delete_removes_exports_and_approval_state(self):
        self.approve()
        self.assertEqual(self.client.get(self.url + '/exports/docx').status_code, 200)
        self.assertEqual(self.client.delete(self.url).status_code, 204)
        self.assertFalse((self.root / self.mid).exists())
        with connect(self.service.repository.database) as db:
            self.assertEqual(db.execute('SELECT count(*) FROM exports').fetchone()[0], 0)
